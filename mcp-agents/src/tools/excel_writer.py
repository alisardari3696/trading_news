import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from src.config import OUTPUT_DIR


def create_excel(all_results: pd.DataFrame, best_results: pd.DataFrame,
                 output_path: Path = None) -> Path:
    if output_path is None:
        OUTPUT_DIR.mkdir(exist_ok=True)
        output_path = OUTPUT_DIR / "mcp_analysis.xlsx"

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        if not best_results.empty:
            best_results.to_excel(writer, sheet_name="Best_Strategies", index=False)
        if not all_results.empty:
            all_results.to_excel(writer, sheet_name="All_Results", index=False)

    _format_workbook(output_path)
    return output_path


def _format_workbook(path: Path):
    wb = load_workbook(path)

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for ws in wb.worksheets:
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx in range(2, ws.max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="center")

        for col_idx in range(1, ws.max_column + 1):
            max_length = 0
            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max(max_length + 2, 10), 25)
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

        ws.freeze_panes = "A2"

    wb.save(path)
