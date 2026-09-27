import numpy as np
import pandas as pd
from pathlib import Path
from zoneinfo import ZoneInfo
import re
import time

DATA_DIR = Path(".")
OUTPUT_DIR = DATA_DIR / "results"

ATR_PERIOD = 14
ATR_TIMEFRAME = "1h"
ATR_SMOOTHING = "wilder"
ATR_MIN_MULTIPLIER = 1.0
ATR_MAX_MULTIPLIER = 6.0
ATR_INTERVAL = 0.5
ATR_MULTIPLIERS = [
    round(ATR_MIN_MULTIPLIER + i * ATR_INTERVAL, 4)
    for i in range(int(round((ATR_MAX_MULTIPLIER - ATR_MIN_MULTIPLIER) / ATR_INTERVAL)) + 1)
]

SAME_CANDLE_RULE = "sl_first"
ESTIMATED_ROLLOVER_FEE_PERCENT_PER_DAY = 0.01
ROLLOVER_HOUR_UTC = 0

MET_PNL = 0
MET_COUNT = 1
MET_WINS = 2
MET_TOTAL = 3

ASSET_TIMEZONES = {
    "USD": "America/New_York",
    "GBP": "Europe/London",
    "EUR": "Europe/Berlin",
    "JPY": "Asia/Tokyo",
    "AUD": "Australia/Sydney",
    "NZD": "Pacific/Auckland",
    "CHF": "Europe/Zurich",
    "CAD": "America/Toronto",
}

pair_sets = {
    "USD": ["EURUSD", "GBPUSD", "USDJPY", "AUDUSD", "NZDUSD", "USDCAD", "USDCHF"],
    "EUR": ["EURUSD", "EURJPY", "EURGBP", "EURAUD", "EURNZD", "EURCHF", "EURCAD"],
    "GBP": ["GBPUSD", "GBPJPY", "EURGBP", "GBPAUD", "GBPNZD", "GBPCHF", "GBPCAD"],
    "JPY": ["USDJPY", "EURJPY", "GBPJPY", "AUDJPY", "NZDJPY", "CADJPY", "CHFJPY"],
    "AUD": ["AUDUSD", "AUDJPY", "EURAUD", "GBPAUD", "AUDNZD", "AUDCHF", "AUDCAD"],
    "NZD": ["NZDUSD", "NZDJPY", "EURNZD", "GBPNZD", "AUDNZD", "NZDCHF", "NZDCAD"],
    "CHF": ["USDCHF", "EURCHF", "GBPCHF", "AUDCHF", "NZDCHF", "CADCHF", "CHFJPY"],
    "CAD": ["USDCAD", "EURCAD", "GBPCAD", "AUDCAD", "NZDCAD", "CADCHF", "CADJPY"],
}


def choose_trading_mode():
    print("\nSelect Trading Mode:")
    print("1. Surprise (Actual vs Forecast)")
    print("2. Candle Colour (Tests Trend and Fade)")
    choice = input("Choice (1/2): ").strip()
    if choice == "1":
        return "surprise"
    if choice == "2":
        return "candle_colour"
    print("Invalid trading mode.")
    return None


def choose_currency_group():
    print("\nAvailable groups:", ", ".join(pair_sets.keys()))
    choice = input("Type currency group (example: GBP): ").upper().strip()
    if choice in pair_sets:
        return choice
    print("Invalid currency group.")
    return None


def choose_entry_time():
    return input("Enter trade entry time in local event timezone HH:MM: ").strip()


def choose_inverted_event():
    print("\nIs higher Actual worse for the currency? (y/n)")
    choice = input("Choice: ").lower().strip()
    return choice == "y"


def choose_grid_settings():
    print("\n--- Window Grid Search Settings ---")
    min_train = int(input("Min Train Window Size (default 3): ").strip() or 3)
    max_train = int(input("Max Train Window Size (default 13): ").strip() or 13)

    min_fwd = int(input("Min Forward Test Window Size (default 3): ").strip() or 3)
    max_fwd = int(input("Max Forward Test Window Size (default 13): ").strip() or 13)

    cutoff_str = input("Top Lowest Delta Percentage Cutoff (e.g., 15 or 30) (default 50): ").strip()
    top_pct = float(cutoff_str) / 100.0 if cutoff_str.replace(".", "", 1).isdigit() else 0.50

    min_wr_str = input("Minimum Acceptable Win Rate % (e.g., 50) (default 70): ").strip()
    min_win_rate = float(min_wr_str) if min_wr_str.replace(".", "", 1).isdigit() else 70.0

    return min_train, max_train, min_fwd, max_fwd, top_pct, min_win_rate


def add_atr_column(df, period=ATR_PERIOD):
    out = df.copy()
    prev_close = out["Close"].shift(1)
    true_range = pd.concat(
        [
            out["High"] - out["Low"],
            (out["High"] - prev_close).abs(),
            (out["Low"] - prev_close).abs(),
        ],
        axis=1,
    ).max(axis=1)
    if ATR_SMOOTHING == "sma":
        out["ATR"] = true_range.rolling(period, min_periods=period).mean()
    else:
        out["ATR"] = true_range.ewm(alpha=1.0 / period, adjust=False, min_periods=period).mean()
    return out


def load_csv(pair):
    file_path = DATA_DIR / f"{pair}60.csv"
    if not file_path.exists():
        print(f"Missing CSV for {pair}: {file_path}")
        return None
    df = pd.read_csv(file_path, sep="\t", header=None, names=["Datetime", "Open", "High", "Low", "Close", "Volume"], engine="c")
    required_cols = ["Datetime", "Open", "High", "Low", "Close"]
    missing_cols = [col for col in required_cols if col not in df.columns]
    if missing_cols:
        print(f"{pair} missing columns: {missing_cols}")
        return None
    df["Datetime"] = pd.to_datetime(df["Datetime"], utc=True, errors="coerce").dt.tz_localize(None)
    for col in ["Open", "High", "Low", "Close"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df = (
        df.dropna(subset=required_cols)
        .sort_values("Datetime")
        .drop_duplicates("Datetime")
        .set_index("Datetime")
    )
    return add_atr_column(df)


def build_pair_arrays(pair, df):
    return {
        "pair": pair,
        "index": df.index,
        "open": df["Open"].to_numpy(),
        "high": df["High"].to_numpy(),
        "low": df["Low"].to_numpy(),
        "close": df["Close"].to_numpy(),
        "atr": df["ATR"].to_numpy(),
    }


def index_position(index, timestamp):
    pos = index.searchsorted(timestamp, side="left")
    if pos < len(index) and index[pos] == timestamp:
        return int(pos)
    return -1


def has_candle(pair_arrays, timestamp):
    return index_position(pair_arrays["index"], timestamp) >= 0


def get_direction(mode, pair, surprise, event_currency, candle=None, sub_mode=None):
    if mode == "surprise":
        if surprise == "neutral":
            return None
        if pair.startswith(event_currency):
            return "long" if surprise == "positive" else "short"
        if pair.endswith(event_currency):
            return "short" if surprise == "positive" else "long"
    elif mode == "candle_colour":
        if candle is None:
            return None
        if candle["Close"] > candle["Open"]:
            candle_colour = "green"
        elif candle["Close"] < candle["Open"]:
            candle_colour = "red"
        else:
            return None
        if sub_mode == "trend":
            return "long" if candle_colour == "green" else "short"
        if sub_mode == "fade":
            return "short" if candle_colour == "green" else "long"
    return None


def paste_news_data(event_currency, entry_time_str, inverted_event, mode):
    print("\n--- Paste News Data ---")
    print("Paste your tab-separated news data below (with header row).")
    print("When done pasting, enter an empty line:\n")
    lines = []
    while True:
        line = input()
        if line.strip() == "" and lines:
            break
        lines.append(line)
    if not lines:
        print("No data pasted.")
        return None
    from io import StringIO
    raw = "\n".join(lines)
    df = pd.read_csv(StringIO(raw), sep="\t")
    if "History" not in df.columns:
        print("Missing 'History' column.")
        return None
    try:
        hour, minute = map(int, entry_time_str.split(":"))
    except ValueError:
        print("Invalid time format. Use HH:MM.")
        return None

    def parse_history_date(date_str):
        date_str = str(date_str).strip()
        date_str = re.sub(r'(\w+\s+\d+)-(\d+),\s*(\d{4})', r'\1, \3', date_str)
        return pd.to_datetime(date_str, format="%b %d, %Y", errors="coerce")

    df["History"] = df["History"].astype(str).apply(parse_history_date)

    if mode == "surprise":
        if "Actual" not in df.columns or "Forecast" not in df.columns:
            print("Surprise mode requires Actual and Forecast columns.")
            return None

        def parse_value(val):
            val = str(val).replace("%", "").replace(",", "").strip().lower()
            if val.endswith("k") or val.endswith("m"):
                return val[:-1]
            return val

        for col in ["Actual", "Forecast"]:
            df[col] = pd.to_numeric(df[col].astype(str).apply(parse_value), errors="coerce")
        df = df.dropna(subset=["History", "Actual", "Forecast"]).copy()
    else:
        df = df.dropna(subset=["History"]).copy()

    local_tz = ZoneInfo(ASSET_TIMEZONES[event_currency])
    utc_tz = ZoneInfo("UTC")

    df["UTC_Time"] = [
        pd.Timestamp(
            year=d.year, month=d.month, day=d.day,
            hour=hour, minute=minute, tz=local_tz,
        ).tz_convert(utc_tz).tz_localize(None)
        for d in df["History"]
    ]

    if mode == "surprise":
        adjusted_surprise = (df["Actual"] - df["Forecast"]) * (-1 if inverted_event else 1)
        df["Surprise_Type"] = adjusted_surprise.apply(
            lambda value: "positive" if value > 0 else "negative" if value < 0 else "neutral"
        )
    else:
        df["Surprise_Type"] = "neutral"

    return df.sort_values("UTC_Time").reset_index(drop=True)


def calculate_pnl(entry_price, exit_price, direction):
    if direction == "long":
        return (exit_price - entry_price) / entry_price
    return (entry_price - exit_price) / entry_price


def calculate_rollover_fee_percent(entry_time, exit_time):
    entry_day = (entry_time - pd.Timedelta(hours=ROLLOVER_HOUR_UTC)).date()
    exit_day = (exit_time - pd.Timedelta(hours=ROLLOVER_HOUR_UTC)).date()
    rollover_count = max(0, (exit_day - entry_day).days)
    return rollover_count * ESTIMATED_ROLLOVER_FEE_PERCENT_PER_DAY


def build_trade_result(entry_price, exit_price, direction, entry_time, exit_time, tp_distance_percent, sl_distance_percent):
    hold_hours = (exit_time - entry_time).total_seconds() / 3600
    gross_pnl_pct = calculate_pnl(entry_price, exit_price, direction)
    rollover_fee_percent = calculate_rollover_fee_percent(entry_time, exit_time)
    rollover_fee_pct = rollover_fee_percent / 100
    return {
        "pnl_pct": gross_pnl_pct - rollover_fee_pct,
        "gross_pnl_pct": gross_pnl_pct,
        "rollover_fee_percent": rollover_fee_percent,
        "hold_hours": hold_hours,
        "entry_price": entry_price,
        "exit_price": exit_price,
        "tp_distance_percent": tp_distance_percent,
        "sl_distance_percent": sl_distance_percent,
    }


def get_end_of_week_cutoff(entry_time):
    days_ahead = 4 - entry_time.weekday()
    if days_ahead < 0 or (days_ahead == 0 and entry_time.hour >= 18):
        days_ahead += 7
    target = entry_time + pd.Timedelta(days=days_ahead)
    return target.replace(hour=18, minute=0, second=0, microsecond=0)


def build_trade_window(pair_arrays, entry_time):
    index = pair_arrays["index"]
    pos = index_position(index, entry_time)
    if pos <= 0:
        return None
    atr = float(pair_arrays["atr"][pos - 1])
    if not np.isfinite(atr) or atr <= 0.0:
        return None
    cutoff = get_end_of_week_cutoff(entry_time)
    end = index.searchsorted(cutoff, side="right")
    if end <= pos:
        return None
    return {
        "entry_time": entry_time,
        "entry_price": float(pair_arrays["open"][pos]),
        "atr": atr,
        "times": index[pos:end],
        "highs": pair_arrays["high"][pos:end],
        "lows": pair_arrays["low"][pos:end],
        "closes": pair_arrays["close"][pos:end],
    }


def simulate_trade(window, direction, tp_atr, sl_atr):
    entry_price = window["entry_price"]
    atr = window["atr"]

    if direction == "long":
        tp_price = entry_price + tp_atr * atr
        sl_price = entry_price - sl_atr * atr
        tp_hits = window["highs"] >= tp_price
        sl_hits = window["lows"] <= sl_price
    else:
        tp_price = entry_price - tp_atr * atr
        sl_price = entry_price + sl_atr * atr
        tp_hits = window["lows"] <= tp_price
        sl_hits = window["highs"] >= sl_price

    tp_distance_percent = abs(tp_price - entry_price) / entry_price * 100
    sl_distance_percent = abs(entry_price - sl_price) / entry_price * 100

    tp_idx = np.flatnonzero(tp_hits)
    sl_idx = np.flatnonzero(sl_hits)
    first_tp = int(tp_idx[0]) if tp_idx.size else -1
    first_sl = int(sl_idx[0]) if sl_idx.size else -1

    if first_tp < 0 and first_sl < 0:
        res = build_trade_result(
            entry_price, float(window["closes"][-1]), direction,
            window["entry_time"], window["times"][-1],
            tp_distance_percent, sl_distance_percent,
        )
        res["exit_reason"] = "Friday Cutoff"
        return res

    if first_sl < 0 or (0 <= first_tp < first_sl):
        exit_idx, exit_price, reason = first_tp, tp_price, "TP"
    elif first_tp < 0 or first_sl < first_tp:
        exit_idx, exit_price, reason = first_sl, sl_price, "SL"
    elif SAME_CANDLE_RULE == "sl_first":
        exit_idx, exit_price, reason = first_sl, sl_price, "SL (Same Candle)"
    else:
        exit_idx, exit_price, reason = first_tp, tp_price, "TP (Same Candle)"

    res = build_trade_result(
        entry_price, exit_price, direction, window["entry_time"],
        window["times"][exit_idx], tp_distance_percent, sl_distance_percent,
    )
    res["exit_reason"] = reason
    return res


def calculate_max_drawdown(pnl_series):
    if len(pnl_series) == 0:
        return 0.0
    cum_pnl = np.cumsum(pnl_series)
    running_max = np.maximum.accumulate(cum_pnl)
    drawdowns = running_max - cum_pnl
    return float(np.max(drawdowns))


def build_atr_reference(unit_sum, unit_count):
    rows = []
    mean_unit = (unit_sum / unit_count) if unit_count > 0 else 0.0
    for mult in ATR_MULTIPLIERS:
        rows.append({
            "atr_multiplier": mult,
            "atr_period": ATR_PERIOD,
            "atr_timeframe": ATR_TIMEFRAME,
            "avg_distance_percent": mult * mean_unit,
        })
    return pd.DataFrame(rows)


def main():
    mode = choose_trading_mode()
    if not mode:
        return

    group = choose_currency_group()
    if not group:
        return

    entry_time = choose_entry_time()
    inverted_event = choose_inverted_event() if mode == "surprise" else False

    min_train, max_train, min_fwd, max_fwd, top_pct_cutoff, min_acceptable_win_rate = choose_grid_settings()

    print(f"\n--- ATR Grid Search Settings ---")
    print(f"ATR({ATR_PERIOD}) computed on {ATR_TIMEFRAME} candles ({ATR_SMOOTHING} smoothing), read from last closed bar before entry")
    print(f"TP multipliers: {ATR_MULTIPLIERS}")
    print(f"SL multipliers: {ATR_MULTIPLIERS}")

    news_df = paste_news_data(group, entry_time, inverted_event, mode)
    if news_df is None:
        return

    pair_dfs = {}
    pair_arrays_map = {}
    for pair in pair_sets[group]:
        df = load_csv(pair)
        if df is not None:
            pair_dfs[pair] = df
            pair_arrays_map[pair] = build_pair_arrays(pair, df)

    if not pair_dfs:
        print("No valid pair data found.")
        return

    total_pasted = len(news_df)

    price_fit_mask = []
    for _, row in news_df.iterrows():
        entry_time_utc = row["UTC_Time"]
        has_price = any(has_candle(pair_arrays_map[p], entry_time_utc) for p in pair_arrays_map)
        price_fit_mask.append(has_price)
    news_df = news_df[price_fit_mask].reset_index(drop=True)
    out_of_range = total_pasted - len(news_df)
    print(f"\nPasted data rows: {total_pasted}")
    print(f"Out of price data range: {out_of_range}")
    print(f"Events matching price data: {len(news_df)}")

    if mode == "surprise":
        news_df = news_df[news_df["Surprise_Type"] != "neutral"].reset_index(drop=True)
        print(f"Filtered out neutral surprise events. Usable events remaining: {len(news_df)}")
    elif mode == "candle_colour":
        valid_mask = []
        for _, row in news_df.iterrows():
            entry_time_utc = row["UTC_Time"]
            prev_time_utc = entry_time_utc - pd.Timedelta(hours=1)
            is_valid = any(
                has_candle(pair_arrays_map[p], entry_time_utc) and has_candle(pair_arrays_map[p], prev_time_utc)
                for p in pair_arrays_map
            )
            valid_mask.append(is_valid)
        news_df = news_df[valid_mask].reset_index(drop=True)
        print(f"Filtered out events with no valid candle. Usable events remaining: {len(news_df)}")

    total_events = len(news_df)
    if total_events == 0:
        print("No usable events.")
        return

    sub_modes = ["trend", "fade"] if mode == "candle_colour" else ["N/A"]

    candidate_keys = []
    for pair in pair_dfs:
        for sub_mode in sub_modes:
            for tp in ATR_MULTIPLIERS:
                for sl in ATR_MULTIPLIERS:
                    candidate_keys.append((pair, sub_mode, tp, sl))
    candidate_pos = {key: i for i, key in enumerate(candidate_keys)}
    cand_tp = np.array([key[2] for key in candidate_keys], dtype=float)
    cand_sl = np.array([key[3] for key in candidate_keys], dtype=float)
    cand_rr = cand_tp / cand_sl
    n_candidates = len(candidate_keys)

    print(f"\nPre-computing ATR trade simulation lookup table ({n_candidates:,} parameter combinations x {total_events} events)...")

    metrics = np.zeros((MET_TOTAL, n_candidates, total_events), dtype=float)
    trade_lookup = {}
    unit_sum = 0.0
    unit_count = 0
    news_rows = news_df.to_dict("records")

    for event_idx, nr in enumerate(news_rows):
        entry_time_utc = nr["UTC_Time"]
        surprise_type = nr["Surprise_Type"]

        for pair, pair_arrays in pair_arrays_map.items():
            window = build_trade_window(pair_arrays, entry_time_utc)
            if window is None:
                continue

            unit_sum += window["atr"] / window["entry_price"] * 100
            unit_count += 1

            if mode == "candle_colour":
                prev_pos = index_position(pair_arrays["index"], entry_time_utc - pd.Timedelta(hours=1))
                if prev_pos < 0:
                    continue
                candle = {
                    "Open": float(pair_arrays["open"][prev_pos]),
                    "Close": float(pair_arrays["close"][prev_pos]),
                }
            else:
                candle = None

            for sub_mode in sub_modes:
                direction = get_direction(
                    mode=mode, pair=pair, surprise=surprise_type,
                    event_currency=group, candle=candle,
                    sub_mode=sub_mode if sub_mode != "N/A" else None,
                )
                if direction is None:
                    continue

                for tp in ATR_MULTIPLIERS:
                    for sl in ATR_MULTIPLIERS:
                        res = simulate_trade(window, direction, tp, sl)
                        if res is None:
                            continue
                        ci = candidate_pos[(pair, sub_mode, tp, sl)]
                        metrics[MET_PNL, ci, event_idx] = res["pnl_pct"] * 100
                        metrics[MET_COUNT, ci, event_idx] = 1.0
                        if res["pnl_pct"] > 0:
                            metrics[MET_WINS, ci, event_idx] = 1.0
                        trade_lookup[(event_idx, ci)] = (
                            res["pnl_pct"] * 100,
                            res["exit_reason"],
                            res["tp_distance_percent"],
                            res["sl_distance_percent"],
                            res["hold_hours"],
                        )

    simulated_combinations = int(metrics[MET_COUNT].sum())
    print(f"Pre-computation complete! Total simulated trades cached: {simulated_combinations:,}")

    prefix = np.empty((MET_TOTAL, n_candidates, total_events + 1), dtype=float)
    prefix[:, :, 0] = 0.0
    np.cumsum(metrics, axis=2, out=prefix[:, :, 1:])
    del metrics

    pre_pnl = prefix[MET_PNL]
    pre_n = prefix[MET_COUNT]
    pre_w = prefix[MET_WINS]
    zeros_c = np.zeros(n_candidates, dtype=float)

    grid_results = []
    best_overall_log = []
    best_overall_pnl = -99999.0
    best_overall_window = None

    print("\nStarting Window Grid Search over specified bounds...")

    for train_size in range(min_train, max_train + 1):
        for fwd_size in range(min_fwd, max_fwd + 1):
            total_window = train_size + fwd_size
            min_req = total_window + 1

            if total_events < min_req:
                continue

            num_iterations = total_events - total_window
            current_trade_logs = []
            skipped_count = 0

            for i in range(num_iterations):
                train_start, train_end = i, i + train_size
                fwd_start, fwd_end = i + train_size, i + total_window
                target_idx = i + total_window

                target_event = news_df.iloc[target_idx]
                target_time_utc = target_event["UTC_Time"]
                history_date = target_event["History"]

                n_tr = pre_n[:, train_end] - pre_n[:, train_start]
                n_fw = pre_n[:, fwd_end] - pre_n[:, fwd_start]
                pool = np.flatnonzero((n_tr > 0) | (n_fw > 0))

                if pool.size == 0:
                    skipped_count += 1
                    current_trade_logs.append({
                        "Iteration": i + 1,
                        "Date": history_date.strftime("%Y-%m-%d"),
                        "Pair": "None",
                        "PnL_Percent": 0.0,
                        "Is_Win": 0,
                        "Reason": "No candidate strategies"
                    })
                    continue

                tr_pnl = pre_pnl[:, train_end] - pre_pnl[:, train_start]
                fw_pnl = pre_pnl[:, fwd_end] - pre_pnl[:, fwd_start]
                tr_avg = np.divide(tr_pnl, n_tr, out=zeros_c.copy(), where=n_tr > 0)
                fw_avg = np.divide(fw_pnl, n_fw, out=zeros_c.copy(), where=n_fw > 0)
                tr_wr = np.divide(pre_w[:, train_end] - pre_w[:, train_start], n_tr, out=zeros_c.copy(), where=n_tr > 0) * 100
                fw_wr = np.divide(pre_w[:, fwd_end] - pre_w[:, fwd_start], n_fw, out=zeros_c.copy(), where=n_fw > 0) * 100

                avg_pnl_delta = np.abs(tr_avg - fw_avg)
                cutoff_count = max(1, int(np.ceil(pool.size * top_pct_cutoff)))
                by_delta = pool[np.argsort(avg_pnl_delta[pool], kind="stable")][:cutoff_count]
                ordered = by_delta[np.lexsort((-cand_rr[by_delta], -fw_avg[by_delta]))]

                acceptable = (
                    (tr_wr >= min_acceptable_win_rate)
                    & (fw_wr >= min_acceptable_win_rate)
                    & (tr_pnl > 0)
                    & (fw_pnl > 0)
                )

                picked = -1
                if acceptable[ordered].any():
                    picked = int(ordered[int(np.argmax(acceptable[ordered]))])
                else:
                    fallback = pool[np.lexsort((-cand_rr[pool], -fw_avg[pool]))]
                    if acceptable[fallback].any():
                        picked = int(fallback[int(np.argmax(acceptable[fallback]))])

                if picked < 0:
                    skipped_count += 1
                    current_trade_logs.append({
                        "Iteration": i + 1,
                        "Date": history_date.strftime("%Y-%m-%d"),
                        "Pair": "None",
                        "PnL_Percent": 0.0,
                        "Is_Win": 0,
                        "Reason": "No strategy with positive train & fwd PnL"
                    })
                    continue

                pair, sub_mode, tp, sl = candidate_keys[picked]
                cached = trade_lookup.get((target_idx, picked))
                if cached is None:
                    skipped_count += 1
                    current_trade_logs.append({
                        "Iteration": i + 1,
                        "Date": history_date.strftime("%Y-%m-%d"),
                        "Pair": "None",
                        "PnL_Percent": 0.0,
                        "Is_Win": 0,
                        "Reason": "Strategy outcome not cached"
                    })
                    continue

                pnl_val, exit_reason, tp_pct, sl_pct, hold_hours = cached
                current_trade_logs.append({
                    "Iteration": i + 1,
                    "Date": history_date.strftime("%Y-%m-%d"),
                    "Pair": pair,
                    "Sub_Mode": sub_mode,
                    "TP_ATR": tp,
                    "SL_ATR": sl,
                    "TP_Percent": tp_pct,
                    "SL_Percent": sl_pct,
                    "Hold_Hours": hold_hours,
                    "PnL_Percent": pnl_val,
                    "Is_Win": 1 if pnl_val > 0 else 0,
                    "Reason": exit_reason
                })

            trades_executed = [tl["PnL_Percent"] for tl in current_trade_logs if tl["Pair"] != "None"]
            total_trades = len(trades_executed)
            actual_total_pnl = sum(trades_executed)
            actual_win_rate = (sum([1 for pnl in trades_executed if pnl > 0]) / total_trades * 100) if total_trades > 0 else 0.0
            actual_ev = (actual_total_pnl / total_trades) if total_trades > 0 else 0.0
            max_dd = calculate_max_drawdown(trades_executed)

            grid_results.append({
                "Train_Size": train_size,
                "Forward_Size": fwd_size,
                "Train_Ratio_Percent": (train_size / total_window) * 100,
                "Total_Window": total_window,
                "Total_Trades": total_trades,
                "Skipped_Trades": skipped_count,
                "Actual_Win_Rate_Percent": actual_win_rate,
                "Actual_Total_PnL_Percent": actual_total_pnl,
                "Actual_EV_Percent": actual_ev,
                "Max_Drawdown_Percent": max_dd
            })

            if actual_total_pnl > best_overall_pnl:
                best_overall_pnl = actual_total_pnl
                best_overall_log = current_trade_logs
                best_overall_window = (train_size, fwd_size)

    if not grid_results:
        print("No valid grid results produced.")
        return

    grid_df = pd.DataFrame(grid_results).sort_values("Actual_Total_PnL_Percent", ascending=False)
    log_df = pd.DataFrame(best_overall_log)
    heatmap_df = pd.DataFrame(grid_results).pivot(index="Train_Size", columns="Forward_Size", values="Actual_Total_PnL_Percent")
    atr_reference_df = build_atr_reference(unit_sum, unit_count)

    OUTPUT_DIR.mkdir(exist_ok=True)
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    out_name = f"{group}_{timestamp}_atr_window_grid_search.xlsx"
    out_path = OUTPUT_DIR / out_name

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        grid_df.to_excel(writer, sheet_name="Window_Grid_Summary", index=False)
        heatmap_df.to_excel(writer, sheet_name="PnL_Heatmap")
        log_df.to_excel(writer, sheet_name="Best_Window_Trade_Log", index=False)
        atr_reference_df.to_excel(writer, sheet_name="ATR_Grid_Reference", index=False)

        from openpyxl.formatting.rule import ColorScaleRule
        ws = writer.sheets["PnL_Heatmap"]
        max_col = ws.max_column
        max_row = ws.max_row
        import openpyxl.utils
        col_letter = openpyxl.utils.get_column_letter(max_col)
        cell_range = f"B2:{col_letter}{max_row}"

        ws.conditional_formatting.add(
            cell_range,
            ColorScaleRule(
                start_type='num', start_value=-10, start_color='F8696B',
                mid_type='num', mid_value=0, mid_color='FFFFFF',
                end_type='num', end_value=10, end_color='63BE7B'
            )
        )

    print(f"\n--- GRID SEARCH FINISHED ---")
    print(f"File saved: {out_path}")
    print(f"Best Window configuration found: {best_overall_window} (Train/Forward)")
    print(f"Total entries in grid: {len(grid_df)}")
    total_skipped = sum(r["Skipped_Trades"] for r in grid_results)
    print(f"Total trades skipped (no positive PnL strategy): {total_skipped}")
    print(f"\nATR reference (average realized distance per multiplier, {total_events} events x {len(pair_dfs)} pairs):")
    print(atr_reference_df.to_string(index=False))


if __name__ == "__main__":
    main()
